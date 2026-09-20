import io
import json
import pytest
import openpyxl
import xlwt
from fastapi.testclient import TestClient

from app.main import app
from app import store
from app.core.database import SessionLocal
from app.core.security import create_access_token, hash_password
from app.core.seed import seed_default_users_and_cpses
from app.models.user import User
from app.services.ingestion.normalizer import normalize_field_name, normalize_record_keys
from app.services.ingestion.detector import detect_file_type
from app.services.ingestion.service import parse_legacy_file

client = TestClient(app)


def auth_header(email: str = "steward@mira.gov.in", role: str = "data_steward") -> dict[str, str]:
    db = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).first()
        if not user:
            user = User(
                email=email,
                password_hash=hash_password("Pass@123"),
                full_name=email,
                role=role,
                is_active=True,
            )
            db.add(user)
            db.commit()
            db.refresh(user)
        token = create_access_token(user)
        return {"Authorization": f"Bearer {token}"}
    finally:
        db.close()


@pytest.fixture(autouse=True)
def setup_state():
    seed_default_users_and_cpses()
    store.reset_stores()
    yield
    store.reset_stores()


# =====================================================================
# 1. FIELD & DETECTOR UNIT TESTS
# =====================================================================

def test_detector_supported_formats():
    assert detect_file_type("data.csv") == "csv"
    assert detect_file_type("materials.txt") == "txt"
    assert detect_file_type("catalog.xml") == "xml"
    assert detect_file_type("items.json") == "json"
    assert detect_file_type("inventory.xls") == "excel"
    assert detect_file_type("master.xlsx") == "excel"
    assert detect_file_type("path/to/FILE.XLSX") == "excel"

    with pytest.raises(ValueError, match="Unsupported file format"):
        detect_file_type("archive.zip")

    with pytest.raises(ValueError, match="Unsupported file format"):
        detect_file_type("document.pdf")


def test_normalizer_field_aliases():
    # material_code variations
    assert normalize_field_name("material_code") == "material_code"
    assert normalize_field_name("Material Code") == "material_code"
    assert normalize_field_name("materialCode") == "material_code"
    assert normalize_field_name("MAT_CODE") == "material_code"
    assert normalize_field_name("item-no") == "material_code"
    assert normalize_field_name("Item Number") == "material_code"
    assert normalize_field_name("CODE") == "material_code"

    # description variations
    assert normalize_field_name("description") == "description"
    assert normalize_field_name("Material Description") == "description"
    assert normalize_field_name("material_name") == "description"
    assert normalize_field_name("itemDesc") == "description"
    assert normalize_field_name("DESC") == "description"

    # unit variations
    assert normalize_field_name("unit") == "unit"
    assert normalize_field_name("UOM") == "unit"
    assert normalize_field_name("unit_of_measure") == "unit"
    assert normalize_field_name("Unit of Measure") == "unit"

    # other fields
    assert normalize_field_name("Material Grade") == "material_grade"
    assert normalize_field_name("Manufacturer Part Number") == "manufacturer_part_number"
    assert normalize_field_name("MFR_PN") == "manufacturer_part_number"
    assert normalize_field_name("Category") == "category"
    assert normalize_field_name("Source Org") == "cpse"


def test_normalize_record_keys_preserves_unmapped():
    raw = {
        "Material-Code": "MAT-001",
        "Material Description": "BALL VALVE SS316 150# 2IN",
        "Unit Of Measure": "PCS",
        "Warehouse Location": "Bay 4",
        "Legacy ID": 9948,
    }
    normalized = normalize_record_keys(raw)
    assert normalized["material_code"] == "MAT-001"
    assert normalized["description"] == "BALL VALVE SS316 150# 2IN"
    assert normalized["unit"] == "PCS"
    assert "other_attributes" in normalized
    assert normalized["other_attributes"]["Warehouse Location"] == "Bay 4"


# =====================================================================
# 2. TXT INGESTION TESTS
# =====================================================================

def test_upload_txt_pipe_delimited_positional():
    headers = auth_header()
    txt_content = (
        "MAT001|Stainless Steel Bolt M10|PCS\n"
        "MAT002|Copper Wire 2.5mm|MTR\n"
    )
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("materials.txt", io.BytesIO(txt_content.encode("utf-8")), "text/plain")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["total_materials"] == 2
    assert data["sample"][0]["material_code"] == "MAT001"
    assert data["sample"][0]["description"] == "Stainless Steel Bolt M10"
    assert data["sample"][0]["unit"] == "PCS"
    assert data["sample"][1]["material_code"] == "MAT002"
    assert data["sample"][1]["unit"] == "MTR"


def test_upload_txt_with_header_row():
    headers = auth_header()
    txt_content = (
        "ITEM_NO|MATERIAL_DESCRIPTION|UOM|MANUFACTURER\n"
        "VLV-100|GATE VALVE CS 150 LB 2 IN|NOS|L&T\n"
        "VLV-200|GLOBE VALVE SS316 300 LB 3 IN|NOS|Audco\n"
    )
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("valves.txt", io.BytesIO(txt_content.encode("utf-8")), "text/plain")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["sample"][0]["material_code"] == "VLV-100"
    assert data["sample"][0]["description"] == "GATE VALVE CS 150 LB 2 IN"
    assert data["sample"][0]["unit"] == "NOS"
    assert data["sample"][0]["manufacturer"] == "L&T"


def test_upload_txt_blank_lines_malformed_rows_and_whitespace():
    headers = auth_header()
    txt_content = (
        "\n"
        "   MAT001   |   Stainless Steel Bolt M10   |   PCS   \n"
        "\n"
        "MALFORMED_ROW_NO_PIPES\n"
        "\n"
        "MAT002|Copper Wire 2.5mm|MTR\n"
        "     \n"
    )
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("messy.txt", io.BytesIO(txt_content.encode("utf-8")), "text/plain")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["sample"][0]["material_code"] == "MAT001"
    assert data["sample"][0]["description"] == "Stainless Steel Bolt M10"
    assert data["sample"][0]["unit"] == "PCS"
    assert data["sample"][1]["material_code"] == "MAT002"


# =====================================================================
# 3. JSON INGESTION TESTS
# =====================================================================

def test_upload_json_array_of_objects():
    headers = auth_header()
    records = [
        {"material_code": "MAT001", "description": "Stainless Steel Bolt M10", "unit": "PCS"},
        {"material_code": "MAT002", "description": "Copper Wire 2.5mm", "unit": "MTR", "category": "Electrical"},
    ]
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("materials.json", io.BytesIO(json.dumps(records).encode("utf-8")), "application/json")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["sample"][0]["material_code"] == "MAT001"
    assert data["sample"][1]["category"] == "Electrical"


def test_upload_json_wrapped_object_and_aliases():
    headers = auth_header()
    payload = {
        "materials": [
            {
                "materialCode": "MAT001",
                "materialDescription": "Stainless Steel Bolt M10",
                "uom": "PCS",
                "manufacturer": "ABC Ltd",
                "grade": "SS304",
            },
            {
                "item_no": "MAT002",
                "item_desc": "Copper Wire 2.5mm",
                "measure_unit": "MTR",
            },
        ]
    }
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("wrapped.json", io.BytesIO(json.dumps(payload).encode("utf-8")), "application/json")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["sample"][0]["material_code"] == "MAT001"
    assert data["sample"][0]["description"] == "Stainless Steel Bolt M10"
    assert data["sample"][0]["unit"] == "PCS"
    assert data["sample"][0]["manufacturer"] == "ABC Ltd"
    assert data["sample"][1]["material_code"] == "MAT002"
    assert data["sample"][1]["description"] == "Copper Wire 2.5mm"
    assert data["sample"][1]["unit"] == "MTR"


def test_upload_json_invalid_structure_and_syntax():
    headers = auth_header()
    # Number as top-level JSON
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("invalid.json", io.BytesIO(b"12345"), "application/json")},
        headers=headers,
    )
    assert resp.status_code == 400
    assert "Invalid JSON top-level structure" in resp.json()["detail"]

    # Malformed JSON syntax
    resp2 = client.post(
        "/api/materials/upload",
        files={"file": ("broken.json", io.BytesIO(b"{material_code: broken"), "application/json")},
        headers=headers,
    )
    assert resp2.status_code == 400
    assert "Malformed JSON" in resp2.json()["detail"]


# =====================================================================
# 4. XML INGESTION TESTS
# =====================================================================

def test_upload_xml_standard_structure():
    headers = auth_header()
    xml_content = """<?xml version="1.0" encoding="UTF-8"?>
<materials>
    <material>
        <code>MAT001</code>
        <description>Stainless Steel Bolt M10</description>
        <unit>PCS</unit>
        <manufacturer>ABC Ltd</manufacturer>
        <category>Fastener</category>
    </material>
    <material>
        <material_code>MAT002</material_code>
        <material_description>Copper Wire 2.5mm</material_description>
        <uom>MTR</uom>
    </material>
</materials>"""
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("materials.xml", io.BytesIO(xml_content.encode("utf-8")), "application/xml")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["sample"][0]["material_code"] == "MAT001"
    assert data["sample"][0]["description"] == "Stainless Steel Bolt M10"
    assert data["sample"][0]["unit"] == "PCS"
    assert data["sample"][0]["manufacturer"] == "ABC Ltd"
    assert data["sample"][0]["category"] == "Fastener"
    assert data["sample"][1]["material_code"] == "MAT002"
    assert data["sample"][1]["description"] == "Copper Wire 2.5mm"
    assert data["sample"][1]["unit"] == "MTR"


def test_upload_xml_attributes_and_missing_optional_fields():
    headers = auth_header()
    xml_content = """<?xml version="1.0" encoding="UTF-8"?>
<items>
    <item code="MAT001">
        <description>Stainless Steel Bolt M10</description>
        <unit>PCS</unit>
    </item>
    <item code="MAT002">
        <description>Seamless Pipe 2 IN SCH40</description>
    </item>
</items>"""
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("items.xml", io.BytesIO(xml_content.encode("utf-8")), "application/xml")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["sample"][0]["material_code"] == "MAT001"
    assert data["sample"][1]["material_code"] == "MAT002"
    assert data["sample"][1]["unit"] is None


def test_upload_xml_malformed_syntax():
    headers = auth_header()
    bad_xml = "<materials><material><code>MAT001</material>"
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("bad.xml", io.BytesIO(bad_xml.encode("utf-8")), "application/xml")},
        headers=headers,
    )
    assert resp.status_code == 400
    assert "Malformed XML" in resp.json()["detail"]


# =====================================================================
# 5. EXCEL (XLS & XLSX) INGESTION TESTS
# =====================================================================

def test_upload_xlsx_workbook():
    headers = auth_header()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Materials"
    ws.append(["Material Code", "Material Description", "Unit of Measure", "Manufacturer", "Category"])
    ws.append(["MAT001", "Stainless Steel Bolt M10", "PCS", "ABC Ltd", "Fastener"])
    ws.append(["MAT002", "Copper Wire 2.5mm", "MTR", "XYZ Ltd", "Electrical"])
    # Empty row to test robustness
    ws.append([None, None, None, None, None])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = client.post(
        "/api/materials/upload",
        files={"file": ("inventory.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["sample"][0]["material_code"] == "MAT001"
    assert data["sample"][0]["description"] == "Stainless Steel Bolt M10"
    assert data["sample"][0]["unit"] == "PCS"
    assert data["sample"][0]["manufacturer"] == "ABC Ltd"
    assert data["sample"][0]["category"] == "Fastener"
    assert data["sample"][1]["material_code"] == "MAT002"


def test_upload_xls_legacy_workbook():
    headers = auth_header()
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")

    headers_row = ["Item No", "Description", "UOM", "Manufacturer"]
    for c, h in enumerate(headers_row):
        ws.write(0, c, h)

    data_rows = [
        ["MAT001", "Stainless Steel Bolt M10", "PCS", "ABC Ltd"],
        ["MAT002", "Copper Wire 2.5mm", "MTR", "XYZ Ltd"],
    ]
    for r, row in enumerate(data_rows, start=1):
        for c, val in enumerate(row):
            ws.write(r, c, val)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = client.post(
        "/api/materials/upload",
        files={"file": ("legacy.xls", buf, "application/vnd.ms-excel")},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_ingested"] == 2
    assert data["sample"][0]["material_code"] == "MAT001"
    assert data["sample"][0]["description"] == "Stainless Steel Bolt M10"
    assert data["sample"][0]["unit"] == "PCS"
    assert data["sample"][0]["manufacturer"] == "ABC Ltd"
    assert data["sample"][1]["material_code"] == "MAT002"


def test_upload_unsupported_format_and_empty_file():
    headers = auth_header()
    resp = client.post(
        "/api/materials/upload",
        files={"file": ("manual.docx", io.BytesIO(b"binary docx content"), "application/octet-stream")},
        headers=headers,
    )
    assert resp.status_code == 400
    assert "Unsupported file format" in resp.json()["detail"]


# =====================================================================
# 6. DETERMINISTIC CROSS-FORMAT EQUIVALENCE TEST
# =====================================================================

def test_cross_format_equivalence():
    """
    Verify that equivalent material records represented in:
    - TXT
    - JSON
    - XML
    - XLSX
    - CSV
    all produce the EXACT SAME internal MIRA material record after passing
    through the complete ingestion, normalization, and specification parsing pipeline.
    """
    headers = auth_header()

    # 1. TXT format
    store.reset_stores()
    txt_content = "MAT001|Stainless Steel Bolt M10|PCS"
    resp_txt = client.post(
        "/api/materials/upload",
        files={"file": ("materials.txt", io.BytesIO(txt_content.encode("utf-8")), "text/plain")},
        headers=headers,
    )
    assert resp_txt.status_code == 200
    assert len(store.MATERIALS) == 1
    record_txt = dict(store.MATERIALS[0])

    # 2. JSON format
    store.reset_stores()
    json_content = json.dumps([{"material_code": "MAT001", "description": "Stainless Steel Bolt M10", "unit": "PCS"}])
    resp_json = client.post(
        "/api/materials/upload",
        files={"file": ("materials.json", io.BytesIO(json_content.encode("utf-8")), "application/json")},
        headers=headers,
    )
    assert resp_json.status_code == 200
    assert len(store.MATERIALS) == 1
    record_json = dict(store.MATERIALS[0])

    # 3. XML format
    store.reset_stores()
    xml_content = "<materials><material><code>MAT001</code><description>Stainless Steel Bolt M10</description><unit>PCS</unit></material></materials>"
    resp_xml = client.post(
        "/api/materials/upload",
        files={"file": ("materials.xml", io.BytesIO(xml_content.encode("utf-8")), "application/xml")},
        headers=headers,
    )
    assert resp_xml.status_code == 200
    assert len(store.MATERIALS) == 1
    record_xml = dict(store.MATERIALS[0])

    # 4. XLSX format
    store.reset_stores()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Material Code", "Description", "Unit"])
    ws.append(["MAT001", "Stainless Steel Bolt M10", "PCS"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp_xlsx = client.post(
        "/api/materials/upload",
        files={"file": ("materials.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp_xlsx.status_code == 200
    assert len(store.MATERIALS) == 1
    record_xlsx = dict(store.MATERIALS[0])

    # 5. CSV format
    store.reset_stores()
    csv_content = "Material Code,Description,Unit\nMAT001,Stainless Steel Bolt M10,PCS\n"
    resp_csv = client.post(
        "/api/materials/upload",
        files={"file": ("materials.csv", io.BytesIO(csv_content.encode("utf-8")), "text/csv")},
        headers=headers,
    )
    assert resp_csv.status_code == 200
    assert len(store.MATERIALS) == 1
    record_csv = dict(store.MATERIALS[0])

    # Verify identical downstream representations across all 5 formats
    canonical_keys = [
        "cpse",
        "material_code",
        "description",
        "normalized_description",
        "category",
        "unit",
        "material_grade",
        "parsed_specifications",
    ]

    for key in canonical_keys:
        assert record_txt[key] == record_csv[key], f"TXT vs CSV mismatch on key '{key}': {record_txt[key]} != {record_csv[key]}"
        assert record_json[key] == record_csv[key], f"JSON vs CSV mismatch on key '{key}': {record_json[key]} != {record_csv[key]}"
        assert record_xml[key] == record_csv[key], f"XML vs CSV mismatch on key '{key}': {record_xml[key]} != {record_csv[key]}"
        assert record_xlsx[key] == record_csv[key], f"XLSX vs CSV mismatch on key '{key}': {record_xlsx[key]} != {record_csv[key]}"

    # Verify content values
    assert record_csv["material_code"] == "MAT001"
    assert record_csv["description"] == "Stainless Steel Bolt M10"
    assert record_csv["normalized_description"] == "STAINLESS STEEL BOLT M10"
    assert record_csv["unit"] == "PCS"
    assert record_csv["category"] == "General"
    assert record_csv["cpse"] == "CPSE_GENERIC"
