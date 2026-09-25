import io
from pathlib import Path
from typing import Any

from app.services.ingestion.normalizer import normalize_record_keys


def _format_cell_value(val: Any) -> Any:
    if val is None:
        return None
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    val_str = str(val).strip()
    return val_str if val_str != "" else None


def parse_excel(content_or_path: bytes | str, filename: str = "") -> list[dict[str, Any]]:
    """
    Unified Excel parser supporting both modern (.xlsx) and legacy (.xls) workbooks.
    Uses openpyxl for .xlsx and xlrd for .xls.
    Normalizes headers via schema field aliases and handles empty rows safely.
    """
    if isinstance(content_or_path, bytes):
        raw_bytes = content_or_path
    elif isinstance(content_or_path, str):
        if not filename:
            filename = content_or_path
        if "\n" not in content_or_path and len(content_or_path) < 1024:
            try:
                with open(content_or_path, "rb") as f:
                    raw_bytes = f.read()
            except (OSError, FileNotFoundError):
                raw_bytes = content_or_path.encode("utf-8")
        else:
            raw_bytes = content_or_path.encode("utf-8")
    else:
        raise ValueError(f"Expected bytes or str for Excel parsing, got {type(content_or_path)}")

    if not raw_bytes:
        return []

    ext = Path(filename).suffix.lower() if filename else ""
    # Detect format via magic bytes if extension not explicit:
    # PK\x03\x04 indicates zip/xlsx, \xD0\xCF\x11\xE0 indicates ole/xls
    is_xlsx = ext == ".xlsx" or raw_bytes.startswith(b"PK\x03\x04")

    all_records: list[dict[str, Any]] = []

    if is_xlsx:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True, read_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_rows: list[list[Any]] = []
                for row in sheet.iter_rows(values_only=True):
                    sheet_rows.append([_format_cell_value(c) for c in row])

                # Find header row in this sheet
                header_idx = -1
                for idx, row in enumerate(sheet_rows):
                    if any(c is not None for c in row):
                        header_idx = idx
                        break

                if header_idx == -1:
                    continue

                headers = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(sheet_rows[header_idx])]
                for row in sheet_rows[header_idx + 1 :]:
                    if not any(c is not None for c in row):
                        continue
                    row_dict: dict[str, Any] = {}
                    for h, v in zip(headers, row):
                        if v is not None:
                            row_dict[h] = v
                    if row_dict:
                        normalized = normalize_record_keys(row_dict)
                        if normalized:
                            normalized["_sheet_name"] = sheet_name
                            all_records.append(normalized)
            wb.close()
        except Exception as exc:
            raise ValueError(f"Malformed XLSX file: {exc}")
    else:
        try:
            import xlrd

            wb = xlrd.open_workbook(file_contents=raw_bytes)
            for sheet_idx in range(wb.nsheets):
                sheet = wb.sheet_by_index(sheet_idx)
                sheet_name = sheet.name
                sheet_rows = []
                for row_idx in range(sheet.nrows):
                    row_vals = [
                        _format_cell_value(sheet.cell_value(row_idx, col_idx))
                        for col_idx in range(sheet.ncols)
                    ]
                    sheet_rows.append(row_vals)

                header_idx = -1
                for idx, row in enumerate(sheet_rows):
                    if any(c is not None for c in row):
                        header_idx = idx
                        break

                if header_idx == -1:
                    continue

                headers = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(sheet_rows[header_idx])]
                for row in sheet_rows[header_idx + 1 :]:
                    if not any(c is not None for c in row):
                        continue
                    row_dict = {}
                    for h, v in zip(headers, row):
                        if v is not None:
                            row_dict[h] = v
                    if row_dict:
                        normalized = normalize_record_keys(row_dict)
                        if normalized:
                            normalized["_sheet_name"] = sheet_name
                            all_records.append(normalized)
        except Exception as exc:
            raise ValueError(f"Malformed XLS file: {exc}")

    return all_records
